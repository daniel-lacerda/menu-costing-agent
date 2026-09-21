"""Load the pantry workbook and derive unit costs by joining the two sheets."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .models import PantryAmendment, PantryItem
from .units import fold, parse_pantry_unit

STOCK_SHEET = "Despensa"
PRICE_SHEET = "Precos"


def load_pantry(path: Path, amendments: list[PantryAmendment] | None = None) -> list[PantryItem]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    stock_rows = list(workbook[STOCK_SHEET].iter_rows(min_row=2, values_only=True))
    price_rows = list(workbook[PRICE_SHEET].iter_rows(min_row=2, values_only=True))
    workbook.close()

    prices = {fold(_text(row[0])): row for row in price_rows if row[0] is not None}
    items: list[PantryItem] = []
    for index, (name, stock, unit_text) in enumerate(
        (row[:3] for row in stock_rows if row[0] is not None), start=2
    ):
        price_row = prices.get(fold(_text(name)))
        if price_row is None:
            raise ValueError(f"{name!r} is in {STOCK_SHEET} but not in {PRICE_SHEET}")
        _, purchased, price_unit, total_paid = price_row[:4]
        if fold(_text(price_unit)) != fold(_text(unit_text)):
            raise ValueError(
                f"{name!r}: unit differs between sheets ({unit_text!r} vs {price_unit!r})"
            )
        unit = parse_pantry_unit(_text(unit_text))
        items.append(
            PantryItem(
                name=_text(name),
                row=index,
                stock=_number(stock),
                unit_text=_text(unit_text),
                base_unit=unit.base,
                per_unit=unit.per_unit,
                purchased=_number(purchased),
                total_paid_brl=_number(total_paid),
            )
        )
    return [_amend(item, amendments or []) for item in items]


def find_item(name: str, items: list[PantryItem]) -> PantryItem | None:
    key = fold(name)
    return next((item for item in items if fold(item.name) == key), None)


def _amend(item: PantryItem, amendments: list[PantryAmendment]) -> PantryItem:
    for amendment in amendments:
        if fold(amendment.name) != fold(item.name):
            continue
        changes: dict[str, object] = {}
        if amendment.package_size is not None and amendment.package_unit is not None:
            changes["base_unit"] = amendment.package_unit
            changes["per_unit"] = amendment.package_size
        if amendment.total_paid_brl is not None:
            changes["total_paid_brl"] = amendment.total_paid_brl
        item = item.model_copy(update=changes)
    return item


def _text(cell: object) -> str:
    if not isinstance(cell, str):
        raise ValueError(f"Expected text in the workbook, got {cell!r}")
    return cell


def _number(cell: object) -> float:
    if isinstance(cell, bool) or not isinstance(cell, int | float):
        raise ValueError(f"Expected a number in the workbook, got {cell!r}")
    return float(cell)
